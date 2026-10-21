import io
import zipfile

import pytest
from conftest import register
from datamind.ingestion import chunk, parse, tabular, validate_upload
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook


def test_chunk_metadata_and_bounds():
    chunks = chunk([{"text": "alpha beta " * 500, "heading": "Terms", "page": 3}])
    assert all(
        len(item["text"]) <= 1200 and item["page"] == 3 and item["heading"] == "Terms" for item in chunks
    )
    assert chunks[0]["offset"] == 0
    assert len(chunks) > 1


@pytest.mark.parametrize(
    "filename,data",
    [
        ("x.exe", b"evil"),
        ("x.pdf", b"not pdf"),
        ("x.docx", b"not zip"),
        ("x.txt", b"\x00"),
        ("x.json", b"not-json"),
        ("x.md", b""),
        ("x.xlsx", b"PKinvalid"),
    ],
)
def test_upload_rejections(filename, data):
    with pytest.raises(HTTPException):
        validate_upload(filename, data)


def test_office_bomb():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        for i in range(3001):
            archive.writestr(str(i), "x")
    with pytest.raises(HTTPException):
        validate_upload("bomb.docx", data.getvalue())


def test_docx_structural_parse(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_heading("Renewal terms", level=1)
    doc.add_paragraph("Customer C001 renews in November.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "C001"
    table.cell(0, 1).text = "30 days"
    doc.save(path)
    blocks = parse(path, ".docx")
    assert any("C001 | 30 days" in block["text"] for block in blocks)
    assert all(block["heading"] == "Renewal terms" for block in blocks)


def test_csv_and_xlsx_actual_duckdb(tmp_path):
    csv = tmp_path / "sales.csv"
    csv.write_text("customer_id,revenue\nC001,120\nC002,80\n")
    profile = tabular(csv, ".csv", tmp_path / "sales.duckdb")
    assert profile["data"]["row_count"] == 2
    assert "INT" in profile["data"]["columns"][1]["type"]
    xlsx = tmp_path / "sales.xlsx"
    book = Workbook()
    book.active.append(["customer_id", "revenue"])
    book.active.append(["C001", 120])
    book.active.append(["C002", 80])
    book.save(xlsx)
    profile = tabular(xlsx, ".xlsx", tmp_path / "excel.duckdb")
    assert profile["data"]["preview"][0]["revenue"] == 120


def test_durable_tabular_job(client):
    from datamind.db import engine
    from datamind.worker import claim, process
    from sqlalchemy.orm import Session

    register(client)
    response = client.post(
        "/api/documents", files={"file": ("orders.csv", b"customer,revenue\nC001,120\nC002,80\n", "text/csv")}
    )
    assert response.status_code == 201, response.text
    source_id = response.json()["id"]
    with Session(engine()) as db:
        job_id = claim(db)
        assert job_id
        assert claim(db) is None
    process(job_id)
    preview = client.get(f"/api/documents/{source_id}").json()
    assert preview["status"] == "ready", preview
    result = client.post(
        f"/api/sources/{source_id}/sql", json={"sql": "SELECT SUM(revenue) AS revenue FROM data"}
    )
    assert result.status_code == 200, result.text
    assert result.json()["rows"] == [{"revenue": 200}]


def test_xlsx_preserves_fractional_amounts(tmp_path):
    path = tmp_path / "amounts.xlsx"
    book = Workbook()
    book.active.append(["amount"])
    book.active.append([120.25])
    book.active.append([-80.75])
    book.save(path)
    profile = tabular(path, ".xlsx", tmp_path / "amounts.duckdb")
    assert [row["amount"] for row in profile["data"]["preview"]] == [120.25, -80.75]
