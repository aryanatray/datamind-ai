import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import duckdb
from docx import Document as WordDocument
from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from datamind.config import settings

EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".json"}


def validate_upload(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in EXTENSIONS or not data or len(data) > settings().max_upload_mb * 1024 * 1024:
        raise HTTPException(400, "Unsupported, empty, or oversized document")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise HTTPException(400, "PDF signature does not match its extension")
    if extension in {".docx", ".xlsx"}:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise HTTPException(400, "Invalid Office document")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if len(archive.infolist()) > 3000 or total > 100 * 1024 * 1024:
                raise HTTPException(400, "Office archive expands beyond safe limits")
            names = set(archive.namelist())
            expected = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
            if expected not in names or any("vbaProject" in name for name in names):
                raise HTTPException(400, "Unsupported Office content")
    if extension in {".txt", ".md", ".csv", ".json"}:
        try:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary content")
            if extension == ".json":
                json.loads(text)
        except (ValueError, UnicodeError) as exc:
            raise HTTPException(400, "Document must contain valid UTF-8 text") from exc
    return extension


def parse(path: Path, extension: str) -> list[dict]:
    if extension == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 1000:
            raise ValueError("Encrypted PDFs and PDFs over 1,000 pages are unsupported")
        pages = [
            {"text": page.extract_text(extraction_mode="layout") or "", "page": i + 1, "heading": ""}
            for i, page in enumerate(reader.pages)
        ]
        if not any(page["text"].strip() for page in pages):
            raise ValueError(
                "No extractable text. OCR is required; convert this scan to a searchable PDF first."
            )
        return pages
    if extension == ".docx":
        doc = WordDocument(path)
        blocks, heading = [], ""
        for item in doc.iter_inner_content():
            if hasattr(item, "text"):
                if item.style and item.style.name.startswith("Heading"):
                    heading = item.text
                if item.text.strip():
                    blocks.append({"text": item.text, "heading": heading, "page": None})
            else:
                text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in item.rows)
                blocks.append({"text": text, "heading": heading, "page": None})
        return blocks
    text = path.read_text(encoding="utf-8-sig")
    if extension == ".json":
        text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    blocks, heading = [], ""
    for paragraph in re.split(r"\n\s*\n", text):
        if paragraph.lstrip().startswith("#"):
            heading = paragraph.splitlines()[0].lstrip("# ")
        if paragraph.strip():
            blocks.append({"text": paragraph.strip(), "heading": heading, "page": None})
    return blocks


def chunk(blocks: list[dict], max_chars: int = 1200, overlap: int = 150) -> list[dict]:
    result = []
    for block in blocks:
        text = block["text"].strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + max_chars)
            if end < len(text):
                boundary = text.rfind(" ", start + max_chars // 2, end)
                if boundary > start:
                    end = boundary
            result.append(
                {**block, "text": text[start:end], "section": block.get("heading", ""), "offset": start}
            )
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    if not result or len(result) > 20000:
        raise ValueError("Document has no text or exceeds the 20,000 chunk limit")
    return result


def identifier(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def tabular(path: Path, extension: str, target: Path) -> dict:
    # Import into an owned DuckDB file before disabling external access on every query connection.
    if target.exists():
        target.unlink()
    with duckdb.connect(str(target), config={"memory_limit": "256MB", "threads": "2"}) as conn:
        if extension == ".csv":
            conn.execute(
                "CREATE TABLE data AS SELECT * FROM read_csv(?, header=true, sample_size=20000)", [str(path)]
            )
        else:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for index, sheet in enumerate(workbook.worksheets[:20]):
                    rows = sheet.iter_rows(values_only=True)
                    header = next(rows, ())
                    if not header:
                        continue
                    if len(header) > 300:
                        raise ValueError("Sheet exceeds 300 columns")
                    names, seen = [], set()
                    for col, value in enumerate(header):
                        name = str(value or f"column_{col + 1}")[:100]
                        while name.lower() in seen:
                            name += "_"
                        seen.add(name.lower())
                        names.append(name)
                    table = "data" if index == 0 else f"sheet_{index + 1}"
                    conn.execute(
                        f"CREATE TABLE {table} ({','.join(identifier(n) + ' VARCHAR' for n in names)})"
                    )
                    batch, count = [], 0
                    for row in rows:
                        batch.append(tuple(None if v is None else str(v) for v in row[: len(names)]))
                        count += 1
                        if count > 250000:
                            raise ValueError("Sheet exceeds 250,000 rows")
                        if len(batch) == 1000:
                            conn.executemany(
                                f"INSERT INTO {table} VALUES ({','.join('?' for _ in names)})", batch
                            )
                            batch = []
                    if batch:
                        conn.executemany(
                            f"INSERT INTO {table} VALUES ({','.join('?' for _ in names)})", batch
                        )
                    # Infer common scalar types without changing ambiguous mixed columns.
                    for name in names:
                        col = identifier(name)
                        for dtype in ("BIGINT", "DOUBLE", "DATE", "BOOLEAN"):
                            if (
                                dtype == "BIGINT"
                                and conn.execute(
                                    f"SELECT count(*) FROM {table} WHERE {col} IS NOT NULL "
                                    f"AND NOT regexp_full_match({col}, '[+-]?[0-9]+')"
                                ).fetchone()[0]
                            ):
                                # DuckDB rounds fractional strings when casting to BIGINT.
                                continue
                            bad = conn.execute(
                                f"SELECT count(*) FROM {table} WHERE {col} IS NOT NULL AND try_cast({col} AS {dtype}) IS NULL"
                            ).fetchone()[0]
                            if bad == 0:
                                conn.execute(
                                    f"ALTER TABLE {table} ALTER COLUMN {col} TYPE {dtype} USING try_cast({col} AS {dtype})"
                                )
                                break
            finally:
                workbook.close()
        tables = [row[0] for row in conn.execute("SHOW TABLES").fetchall()]
        profile = {}
        for table in tables:
            cols = conn.execute(f"DESCRIBE {identifier(table)}").fetchall()
            profile[table] = {
                "columns": [{"name": c[0], "type": c[1]} for c in cols],
                "row_count": conn.execute(f"SELECT count(*) FROM {identifier(table)}").fetchone()[0],
                "null_counts": {
                    c[0]: conn.execute(
                        f"SELECT count(*) FROM {identifier(table)} WHERE {identifier(c[0])} IS NULL"
                    ).fetchone()[0]
                    for c in cols
                },
                "preview": json.loads(
                    json.dumps(
                        [
                            dict(zip([c[0] for c in cols], row, strict=True))
                            for row in conn.execute(f"SELECT * FROM {identifier(table)} LIMIT 20").fetchall()
                        ],
                        default=str,
                    )
                ),
            }
        return profile


def checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
